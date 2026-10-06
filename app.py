import os
import sqlite3
import hashlib
import hmac
import secrets
from datetime import datetime, timezone
from pathlib import Path
from typing import Optional
from io import BytesIO

from fastapi import FastAPI, Depends, HTTPException, Request
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field
from starlette.middleware.sessions import SessionMiddleware
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent
DB_PATH = Path(os.getenv("DB_PATH", str(BASE_DIR / "repair_app.db")))
SESSION_SECRET = os.getenv("SESSION_SECRET", "change-me-in-production")
COOKIE_HTTPS_ONLY = os.getenv("COOKIE_HTTPS_ONLY", "1") == "1"
REQUEST_TYPES = {"diagnostics": "Диагностика", "repair": "Ремонт", "tire": "Шиномонтаж", "maintenance": "ТО"}
STATUSES = {"created": "Создана", "in_progress": "В работе", "waiting_parts": "Ожидаются запчасти", "completed": "Выполнена"}
app = FastAPI(title="Автопарк — заявки на ремонт", version="1.0.0")
app.add_middleware(SessionMiddleware, secret_key=SESSION_SECRET, https_only=COOKIE_HTTPS_ONLY, same_site="lax", session_cookie="repair_session", max_age=60*60*24*7)
app.mount("/static", StaticFiles(directory=str(BASE_DIR / "static")), name="static")

def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn

def now_iso():
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()

def today_iso():
    return datetime.now().astimezone().date().isoformat()

def hash_password(password, salt=None):
    salt = salt or secrets.token_bytes(16)
    dk = hashlib.pbkdf2_hmac("sha256", password.encode(), salt, 210_000)
    return f"pbkdf2_sha256$210000${salt.hex()}${dk.hex()}"

def verify_password(password, stored):
    try:
        algo, rounds, salt_hex, hash_hex = stored.split("$")
        if algo != "pbkdf2_sha256": return False
        candidate = hashlib.pbkdf2_hmac("sha256", password.encode(), bytes.fromhex(salt_hex), int(rounds))
        return hmac.compare_digest(candidate.hex(), hash_hex)
    except Exception:
        return False

def init_db():
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    conn = get_db()
    conn.executescript("""
    CREATE TABLE IF NOT EXISTS users (id INTEGER PRIMARY KEY AUTOINCREMENT, username TEXT UNIQUE NOT NULL, password_hash TEXT NOT NULL, role TEXT NOT NULL CHECK(role IN ('manager','viewer')), created_at TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS repair_requests (id INTEGER PRIMARY KEY AUTOINCREMENT, request_type TEXT NOT NULL, plate TEXT NOT NULL, comment TEXT NOT NULL DEFAULT '', status TEXT NOT NULL, created_at TEXT NOT NULL, updated_at TEXT NOT NULL, completed_at TEXT, created_by TEXT NOT NULL);
    CREATE INDEX IF NOT EXISTS idx_requests_created_at ON repair_requests(created_at);
    CREATE INDEX IF NOT EXISTS idx_requests_status ON repair_requests(status);
    """)
    admin_password = os.getenv("ADMIN_PASSWORD", "admin123")
    viewer_password = os.getenv("VIEWER_PASSWORD", "viewer123")
    if conn.execute("SELECT COUNT(*) c FROM users").fetchone()[0] == 0:
        conn.execute("INSERT INTO users(username,password_hash,role,created_at) VALUES(?,?,?,?)", ("admin", hash_password(admin_password), "manager", now_iso()))
        conn.execute("INSERT INTO users(username,password_hash,role,created_at) VALUES(?,?,?,?)", ("viewer", hash_password(viewer_password), "viewer", now_iso()))
    conn.commit(); conn.close()
init_db()

class LoginIn(BaseModel):
    username: str = Field(min_length=1, max_length=64)
    password: str = Field(min_length=1, max_length=256)
class RequestIn(BaseModel):
    request_type: str
    plate: str = Field(min_length=1, max_length=32)
    comment: str = Field(default="", max_length=1000)
class StatusIn(BaseModel):
    status: str

def current_user(request): return request.session.get("user")
def require_login(request):
    user=current_user(request)
    if not user: raise HTTPException(401,"Требуется вход")
    return user
def require_manager(request):
    user=require_login(request)
    if user.get("role") != "manager": raise HTTPException(403,"Недостаточно прав")
    return user
def require_csrf(request):
    if request.method in {"GET","HEAD","OPTIONS"}: return
    token=request.session.get("csrf")
    header=request.headers.get("X-CSRF-Token")
    if not token or not header or not hmac.compare_digest(token,header): raise HTTPException(403,"CSRF token invalid")
def row_to_dict(r):
    d=dict(r); d["request_type_label"]=REQUEST_TYPES.get(d["request_type"],d["request_type"]); d["status_label"]=STATUSES.get(d["status"],d["status"]); d["created_date"]=d["created_at"][:10]; return d

@app.get("/")
def root(): return FileResponse(BASE_DIR/"static"/"index.html")
@app.post("/api/login")
def login(payload:LoginIn, request:Request):
    conn=get_db(); user=conn.execute("SELECT username,password_hash,role FROM users WHERE username=?",(payload.username.strip(),)).fetchone(); conn.close()
    if not user or not verify_password(payload.password,user["password_hash"]): raise HTTPException(401,"Неверный логин или пароль")
    request.session.clear(); request.session["user"]={"username":user["username"],"role":user["role"]}; request.session["csrf"]=secrets.token_urlsafe(24)
    return {"ok":True,"user":request.session["user"],"csrf":request.session["csrf"]}
@app.post("/api/logout")
def logout(request:Request): require_csrf(request); request.session.clear(); return {"ok":True}
@app.get("/api/me")
def me(request:Request):
    user=current_user(request); return {"authenticated":bool(user),"user":user,"csrf":request.session.get("csrf")}
@app.get("/api/requests")
def list_requests(request:Request, from_date:Optional[str]=None, to_date:Optional[str]=None, active:int=0):
    require_login(request); conn=get_db(); sql="SELECT * FROM repair_requests WHERE 1=1"; args=[]
    if active: sql += " AND status != 'completed'"
    if from_date: sql += " AND date(created_at) >= date(?)"; args.append(from_date)
    if to_date: sql += " AND date(created_at) <= date(?)"; args.append(to_date)
    rows=conn.execute(sql+" ORDER BY datetime(created_at) DESC,id DESC",args).fetchall(); conn.close(); return {"items":[row_to_dict(r) for r in rows]}
@app.post("/api/requests")
def create_request(payload:RequestIn, request:Request, user=Depends(require_manager)):
    require_csrf(request)
    if payload.request_type not in REQUEST_TYPES: raise HTTPException(400,"Неизвестный вид заявки")
    plate=payload.plate.strip().upper()
    if not plate: raise HTTPException(400,"Введите госномер")
    ts=now_iso(); conn=get_db(); cur=conn.execute("INSERT INTO repair_requests(request_type,plate,comment,status,created_at,updated_at,completed_at,created_by) VALUES(?,?,?,?,?,?,?,?)",(payload.request_type,plate,payload.comment.strip(),"created",ts,ts,None,user["username"])); conn.commit(); row=conn.execute("SELECT * FROM repair_requests WHERE id=?",(cur.lastrowid,)).fetchone(); conn.close(); return row_to_dict(row)
@app.patch("/api/requests/{request_id}/status")
def update_status(request_id:int,payload:StatusIn,request:Request,user=Depends(require_manager)):
    require_csrf(request)
    if payload.status not in STATUSES: raise HTTPException(400,"Неизвестный статус")
    ts=now_iso(); conn=get_db(); existing=conn.execute("SELECT * FROM repair_requests WHERE id=?",(request_id,)).fetchone()
    if not existing: conn.close(); raise HTTPException(404,"Заявка не найдена")
    conn.execute("UPDATE repair_requests SET status=?,updated_at=?,completed_at=? WHERE id=?",(payload.status,ts,ts if payload.status=="completed" else None,request_id)); conn.commit(); row=conn.execute("SELECT * FROM repair_requests WHERE id=?",(request_id,)).fetchone(); conn.close(); return row_to_dict(row)
@app.get("/api/stats")
def stats(request:Request,from_date:Optional[str]=None,to_date:Optional[str]=None):
    require_login(request); from_date=from_date or today_iso(); to_date=to_date or from_date; conn=get_db()
    daily=conn.execute("""SELECT date(created_at) day,COUNT(*) total,SUM(CASE WHEN request_type='diagnostics' THEN 1 ELSE 0 END) diagnostics,SUM(CASE WHEN request_type='repair' THEN 1 ELSE 0 END) repair,SUM(CASE WHEN request_type='tire' THEN 1 ELSE 0 END) tire,SUM(CASE WHEN request_type='maintenance' THEN 1 ELSE 0 END) maintenance,SUM(CASE WHEN status='created' THEN 1 ELSE 0 END) created,SUM(CASE WHEN status='in_progress' THEN 1 ELSE 0 END) in_progress,SUM(CASE WHEN status='waiting_parts' THEN 1 ELSE 0 END) waiting_parts,SUM(CASE WHEN status='completed' THEN 1 ELSE 0 END) completed FROM repair_requests WHERE date(created_at) BETWEEN date(?) AND date(?) GROUP BY date(created_at) ORDER BY day DESC""",(from_date,to_date)).fetchall()
    period_total=conn.execute("SELECT COUNT(*) c FROM repair_requests WHERE date(created_at) BETWEEN date(?) AND date(?)",(from_date,to_date)).fetchone()[0]; active=conn.execute("SELECT COUNT(*) c FROM repair_requests WHERE status!='completed'").fetchone()[0]; conn.close()
    return {"daily":[dict(r) for r in daily],"period_total":period_total,"active_total":active}

def build_workbook(from_date,to_date):
    conn=get_db(); requests=conn.execute("SELECT * FROM repair_requests WHERE date(created_at) BETWEEN date(?) AND date(?) ORDER BY datetime(created_at),id",(from_date,to_date)).fetchall(); daily=conn.execute("""SELECT date(created_at) day,COUNT(*) total,SUM(CASE WHEN request_type='diagnostics' THEN 1 ELSE 0 END) diagnostics,SUM(CASE WHEN request_type='repair' THEN 1 ELSE 0 END) repair,SUM(CASE WHEN request_type='tire' THEN 1 ELSE 0 END) tire,SUM(CASE WHEN request_type='maintenance' THEN 1 ELSE 0 END) maintenance,SUM(CASE WHEN status='created' THEN 1 ELSE 0 END) created,SUM(CASE WHEN status='in_progress' THEN 1 ELSE 0 END) in_progress,SUM(CASE WHEN status='waiting_parts' THEN 1 ELSE 0 END) waiting_parts,SUM(CASE WHEN status='completed' THEN 1 ELSE 0 END) completed FROM repair_requests WHERE date(created_at) BETWEEN date(?) AND date(?) GROUP BY date(created_at) ORDER BY day""",(from_date,to_date)).fetchall(); conn.close()
    wb=Workbook(); ws=wb.active; ws.title="Заявки"; ws.append(["ID","Дата","Время","Госномер","Вид заявки","Комментарий","Статус","Создал","Обновлена","Выполнена"])
    for r in requests:
        dt=datetime.fromisoformat(r["created_at"]); upd=datetime.fromisoformat(r["updated_at"]); comp=datetime.fromisoformat(r["completed_at"]) if r["completed_at"] else None
        ws.append([r["id"],dt.date().isoformat(),dt.strftime("%H:%M"),r["plate"],REQUEST_TYPES.get(r["request_type"],r["request_type"]),r["comment"],STATUSES.get(r["status"],r["status"]),r["created_by"],upd.isoformat(),comp.isoformat() if comp else ""])
    ws2=wb.create_sheet("По дням"); ws2.append(["Дата","Всего","Диагностика","Ремонт","Шиномонтаж","ТО","Создана","В работе","Ожидаются запчасти","Выполнена"])
    for r in daily: ws2.append([r["day"],r["total"],r["diagnostics"],r["repair"],r["tire"],r["maintenance"],r["created"],r["in_progress"],r["waiting_parts"],r["completed"]])
    ws3=wb.create_sheet("По видам"); ws3.append(["Вид заявки","Количество"])
    for t in REQUEST_TYPES: ws3.append([REQUEST_TYPES[t],sum(r[t] or 0 for r in daily)])
    ws4=wb.create_sheet("По статусам"); ws4.append(["Статус","Количество"])
    for s in STATUSES: ws4.append([STATUSES[s],sum(r[s] or 0 for r in daily)])
    fill=PatternFill("solid",fgColor="1F4E78"); font=Font(color="FFFFFF",bold=True); thin=Side(style="thin",color="D9E2F3")
    for sheet in wb.worksheets:
        sheet.freeze_panes="A2"
        for c in sheet[1]: c.fill=fill; c.font=font; c.alignment=Alignment(horizontal="center",vertical="center"); c.border=Border(bottom=thin)
        for row in sheet.iter_rows(min_row=2):
            for c in row: c.alignment=Alignment(vertical="top",wrap_text=True)
        for col in range(1,sheet.max_column+1):
            ml=max((len(str(x.value)) if x.value is not None else 0) for row in sheet.iter_rows(min_col=col,max_col=col) for x in row); sheet.column_dimensions[get_column_letter(col)].width=min(max(ml+2,10),42)
    ws.column_dimensions["F"].width=50; ws.auto_filter.ref=ws.dimensions; ws2.auto_filter.ref=ws2.dimensions
    out=BytesIO(); wb.save(out); out.seek(0); return out
@app.get("/api/export.xlsx")
def export_xlsx(request:Request,from_date:Optional[str]=None,to_date:Optional[str]=None):
    require_login(request); from_date=from_date or today_iso(); to_date=to_date or from_date; stream=build_workbook(from_date,to_date); return StreamingResponse(stream,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":f'attachment; filename="zayavki_{from_date}_{to_date}.xlsx"'})
